from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

attendees = [
    ("YU ARAKAWA", "Consultant", "IBM Japan, Ltd."),
    ("Suli Abas", "Business Development Manager", "NSIGHT, INC."),
    ("Katrina Acebedo", "Data Architect", "The Hershey Company"),
    ("NTT DATA Acuna Delaveau", "Directora Enterprise Solution", "NTT DATA Chile"),
    ("Laura Adare", "Manager/Head of Dept", "State of Louisiana"),
    ("Dhruv Aggarwal", "Partner Business Officer to SAP's Chi…", "SAP SE"),
    ("Yeshu Aggarwal", "Principal Value Advisor", "SAP Canada"),
    ("Mukul Agrawal", "IT Director", "Cimpress USA Incorporated"),
    ("Lucia Aguilar Ledezma", "RISE Business Developer", "SAP Costa Rica, S.A."),
    ("Saritha Aguru", "Service Delivery & Product M…", "Ernst & Young"),
    ("Takuya Akatsuka", "GM", "株式会社リクルート"),
    ("Zia Akbar", "Sr Architect", "Jabil Circuit, Inc. c/o Compaq"),
    ("Ram Akella", "Value Advisor", "SAP AMERICA, INC."),
    ("Kash Al Aziz", "Global Head, Corporate Seg…", "SAP MENA LLC"),
    ("Ebenezer Aladeojebi", "Product Manager", "J.P Morgan Chase & Co."),
    ("Torsten Albert", "VP - Head of CPIT Technology Servic…", "SAP SE"),
    ("Steven Albor Jr", "BTP…", "SAP America Public Sec.(Newtown Sq)"),
    ("Cristian Ali", "PES Head MCLAC", "SAP"),
    ("Dilia Alicea", "Management Support", "SAP America (Miami)"),
    ("Benjamin Alimi", "Director", "Panaya Ltd."),
    ("Mathieu Allard", "VP Finance", "Coveo Solutions Inc."),
    ("Kristen Allman", "Sr. Solution Sales Execut…", "SAP AMERICA, INC."),
    ("Salazar del Rio Alma Aurora", "Functional leader", "Bachoco, S.A. de C.V."),
    ("Aditya Amar", "Senior Softw…", "DOW JONES & COMPANY, INC."),
    ("Shafaqat Ameez", "VP Sales - Enterprise Studios", "Globant, LLC"),
    ("Teja Amerineni", "Cheif Strategic officer", "ReleaseOwl"),
    ("David Amiot", "Managing Director", "Deloitte"),
    ("Barbi Amodeo", "Revenue Man…", "TAMKO Building Products, LLC"),
    ("Anu Anand", "Chief of Staff", "SAP INDIA PVT. LTD."),
    ("Piyush Anandani", "Direct…", "Hewlett Packard Enterprise Company"),
    ("Richard Anderson", "SSE - SCM", "SAP AMERICA, INC."),
    ("Tiffany Anderson", "Manager, SA…", "Carlisle Construction Materials"),
    ("Joni Angkasa", "Regional Head of Growth S…", "SAP ASIA Pte Ltd"),
    ("Sean Antonello", "VP, SAP Analytics", "Improving Toronto Inc."),
    ("Karthik Anupoju", "Business Data Analyst", "KATBOTZ LLC"),
    ("Jamie Anushko", "Value Advisor", "SAP AMERICA, INC."),
    ("Kohei Aoki", "Senior Global Account E…", "SAP Japan Co., Ltd."),
    ("Ali Aouzal", "SAP-SCM Consulta…", "Previan Technologies Inc."),
    ("Amy Arendt", "VP, IT Service &…", "Oldcastle Infrastructure, Inc."),
    ("Melissa Arin", "Marketing & Alliances Manager", "GyanSys, Inc."),
    ("Akin Aritmac", "Chief Sales Officer", "sovanta America Inc."),
    ("Dhanishtha Arora", "Consult…", "Celebal Technologies Americas, Inc."),
    ("Tanya Arora", "Global…", "Celebal Technologies Americas Inc."),
    ("Chris Arthur", "Associate Director - Soluti…", "Collins Aerospace"),
    ("Miguel Artigas", "CEO", "724BC Consulting SAS"),
    ("Sokrates Aslanidis", "Territory Account Manager —…", "Releaseowl Inc"),
    ("Michael Asomaning", "TECHNICAL LEAD", "Defense Logistics Agency"),
    ("Da-Wei Au", "Group Product Manag…", "Autodesk Asia Pte Ltd"),
    ("Kevin Aubrey", "Marketing Lead - AMS", "SecurityBridge GmbH"),
    ("Adib Audi", "Account Executive", "SAP Canada"),
    ("Andreas Axer", "", "WIRTGEN GROUP Branch of John Deere Gmb…"),
    ("Marisol Ayala", "VP, Corporate…", "Inspire Medical Systems, Inc."),
    ("Angel Azat", "Executive Assistant", "SAP AMERICA, INC."),
]

wb = Workbook()
ws = wb.active
ws.title = "Attendees"

headers = ["Name", "Role", "Company"]
ws.append(headers)

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="left", vertical="center")

for row in attendees:
    ws.append(row)

ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 42
ws.column_dimensions["C"].width = 42
ws.freeze_panes = "A2"

wb.save("/home/user/test-one/attendees.xlsx")
print(f"Wrote {len(attendees)} attendees to attendees.xlsx")
